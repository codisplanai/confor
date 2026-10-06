"""
Paridade do Excel: compara a planilha gerada pelo webapp (tests/out/Conciliacao_gerada.xlsx)
com a planilha entregue pelo script Python (Conciliacao_Fornecedores_Agosto_2026.xlsx).

Compara, aba a aba: valores, fórmulas, formato numérico, preenchimento e dimensões.
Diferenças esperadas e ignoradas (documentadas no plano):
  - nomes de abas com contagem dinâmica (comparados por prefixo);
  - textos com números dentro (formato pt-BR no webapp, americano no Python);
  - subtítulo (inclui CNPJ) e notas técnicas (agora dinâmicas) do Resumo Executivo.

Uso: node webapp/tests/gen_xlsx.mjs ; python webapp/tests/excel_parity.py
"""
import os
import re
import sys

import openpyxl

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
ORIG = os.path.join(RAIZ, "Conciliacao_Fornecedores_Agosto_2026.xlsx")
NOVO = os.path.join(os.path.dirname(__file__), "out", "Conciliacao_gerada.xlsx")

wo = openpyxl.load_workbook(ORIG)
wn = openpyxl.load_workbook(NOVO)

erros = []


def norm_txt(v):
    """Neutraliza números em textos (formatos diferentes) e espaços; None == ''."""
    if v is None:
        return ""
    if isinstance(v, str):
        return re.sub(r"[-+]?[\d.,]+", "#", v).strip()
    return v


def larguras(ws):
    """Largura por índice de coluna (openpyxl agrupa colunas adjacentes iguais em min/max)."""
    out = {}
    for dim in ws.column_dimensions.values():
        if dim.width and dim.min and dim.max:
            for i in range(dim.min, dim.max + 1):
                out[i] = dim.width
    return out


def num_eq(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 0.005
    return a == b


nomes_o = wo.sheetnames
nomes_n = wn.sheetnames
base = lambda n: re.sub(r"\s*\(\d+\)$", "", n)
if [base(n) for n in nomes_o] != [base(n) for n in nomes_n]:
    erros.append(f"Abas diferentes: {nomes_o} x {nomes_n}")

for so, sn in zip(nomes_o, nomes_n):
    a, b = wo[so], wn[sn]
    resumo = so == "Resumo Executivo"
    print(f"Aba '{so}' -> '{sn}' | linhas {a.max_row}/{b.max_row} colunas {a.max_column}/{b.max_column}")
    if not resumo and (a.max_row != b.max_row or a.max_column != b.max_column):
        erros.append(f"[{so}] dimensões {a.max_row}x{a.max_column} != {b.max_row}x{b.max_column}")
    if a.freeze_panes != b.freeze_panes:
        erros.append(f"[{so}] freeze_panes {a.freeze_panes} != {b.freeze_panes}")

    # larguras de coluna
    la, lb = larguras(a), larguras(b)
    for idx, w in la.items():
        w2 = lb.get(idx)
        if w2 is None or abs(w - w2) > 0.8:
            erros.append(f"[{so}] largura coluna {idx}: {w} != {w2}")

    for row in a.iter_rows():
        for ca in row:
            cb = b[ca.coordinate]
            va, vb = ca.value, cb.value
            if resumo and ca.row in (3,):
                continue  # subtítulo
            if resumo and ca.row > 13 and isinstance(va, str) and re.match(r"^\d\. ", va or ""):
                continue  # notas técnicas dinâmicas
            if isinstance(va, str) and va.startswith("="):
                if va != vb:
                    erros.append(f"[{so}] {ca.coordinate} fórmula {va} != {vb}")
            elif isinstance(va, str) or isinstance(vb, str):
                if norm_txt(va) != norm_txt(vb):
                    erros.append(f"[{so}] {ca.coordinate} texto '{va}' != '{vb}'")
            elif not num_eq(va, vb) and not (va in (None, "") and vb in (None, "")):
                erros.append(f"[{so}] {ca.coordinate} valor {va} != {vb}")
            if va is not None and ca.number_format != cb.number_format and isinstance(va, (int, float)):
                erros.append(f"[{so}] {ca.coordinate} formato {ca.number_format} != {cb.number_format}")
            fa = ca.fill.fgColor.rgb if ca.fill and ca.fill.fill_type else None
            fb = cb.fill.fgColor.rgb if cb.fill and cb.fill.fill_type else None
            norm = lambda c: c[-6:] if isinstance(c, str) else c
            if norm(fa) != norm(fb) and not (resumo and ca.row > 13):
                erros.append(f"[{so}] {ca.coordinate} cor {fa} != {fb}")

    # células mescladas
    ma = sorted(str(m) for m in a.merged_cells.ranges)
    mb = sorted(str(m) for m in b.merged_cells.ranges)
    if ma != mb and not resumo:
        erros.append(f"[{so}] mesclagens {ma[:3]} != {mb[:3]}")

print()
if erros:
    print(f"{len(erros)} diferença(s):")
    for e in erros[:40]:
        print(" -", e)
    sys.exit(1)
print("Paridade do Excel OK: valores, fórmulas, formatos, cores, larguras e mesclagens idênticos.")
